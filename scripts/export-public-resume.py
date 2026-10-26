"""Export the public resume from the same JSON used by React, with the fixed V2 template.
Usage: bundled-python scripts/export-public-resume.py [--output-dir PATH]
No JD matching, network calls, keys, user data store, or production deployment.
"""
import argparse,hashlib,json,os,subprocess,tempfile
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt,Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
R=Path(__file__).resolve().parents[1]
DEP=Path('/Users/bettercallme/.cache/codex-runtimes/codex-primary-runtime/dependencies')
RENDER=Path(os.environ.get('RESUME_RENDERER','/Users/bettercallme/.codex/plugins/cache/openai-primary-runtime/documents/26.1024.11018/skills/documents/render_docx.py'))
def hash_file(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def border(p,edge='bottom'):
 b=OxmlElement('w:pBdr');e=OxmlElement('w:'+edge)
 for k,v in [('val','single'),('sz','10'),('color','000000'),('space','3')]:e.set(qn('w:'+k),v)
 b.append(e);p._p.get_or_add_pPr().append(b)
def text(d,s,bullet=False):
 p=d.add_paragraph();p.paragraph_format.keep_together=True
 if bullet:p.paragraph_format.left_indent=Pt(10);p.paragraph_format.first_line_indent=Pt(-10);p.add_run('•  ')
 p.add_run(s);return p
def heading(d,s):p=d.add_paragraph(s,'Heading 1');border(p)
def item_title(d,title,date):
 p=d.add_paragraph(style='Heading 2');p.paragraph_format.tab_stops.add_tab_stop(Mm(178),2);p.add_run(title);run=p.add_run('\t'+date.replace(' — ','—'));run.bold=False;run.font.size=Pt(10)
def export(out):
 data_path=R/'src/data/resume.json';data=json.loads(data_path.read_text());profile=data['profile'];d=Document(R/'assets/resume/resume_v2.docx')
 names={'{{NAME}}':profile['name'],'{{ROLE}}':profile['direction']+' | 2027届本科','{{CONTACT}}':'邮箱 '+profile['email'],'{{PORTFOLIO}}':'作品集 '+profile['portfolio']}
 for p in d.tables[0].cell(0,0).paragraphs:
  for run in p.runs:
   for k,v in names.items():run.text=run.text.replace(k,v)
 p=d.paragraphs[0];p.text='个人概况  '+profile['summary'];p.runs[0].font.size=Pt(10)
 heading(d,'增长与商业化经历')
 for x in data['experiences']:
  item_title(d,x['title']+'｜'+x['subtitle'],x['period'])
  for b in x['bullets']:
   if not b.get('pdfOmit'):text(d,b.get('pdfText',b['text']),True)
  if x.get('cooperation'):
   c=x['cooperation'];p=text(d,c['title']);p.runs[0].bold=True
   text(d,c['pdfText'])
 heading(d,'AI产品与项目实践')
 for x in data['projects']:
  item_title(d,x['title']+'｜'+x['role'],x['period'])
  for b in x['bullets']:text(d,b.get('pdfText',b['text']),True)
 heading(d,'教育背景与荣誉');e=data['education'];item_title(d,e['school']+'｜'+e['major']+'｜本科',e['period']+'（预计）')
 text(d,' · '.join(e['honors'][:3]));text(d,e['honors'][3])
 heading(d,'专业能力')
 for group in data['skills']:text(d,group['pdfTitle']+'：'+' / '.join(group['items']))
 out.mkdir(parents=True,exist_ok=True);docx=out/'he-yang-resume.docx';d.save(docx)
 config=out/'fonts.conf';config.write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd"><fontconfig><dir>/System/Library/Fonts</dir><dir>/Library/Fonts</dir><cachedir>{out}/font-cache</cachedir></fontconfig>')
 env={**os.environ,'FONTCONFIG_FILE':str(config.resolve())}
 subprocess.run([str(DEP/'python/bin/python3'),str(RENDER),str(docx),'--output_dir',str(out),'--emit_pdf'],check=True,env=env,capture_output=True)
 pdf=out/'he-yang-resume.pdf';reader=PdfReader(pdf);full=''.join(p.extract_text() or '' for p in reader.pages);norm=''.join(full.split())
 assert len(reader.pages)==1,'Public PDF must be one page; shorten redundant copy before changing font'
 assert abs(float(reader.pages[0].mediabox.width)-595.28)<1 and abs(float(reader.pages[0].mediabox.height)-841.89)<1
 for required in [profile['name'],profile['email'],profile['portfolio'],'天津丁尘文化传媒','杭州树成林教育科技集团有限公司','远程合作','10万+','470万+','50+']:
  assert ''.join(required.split()) in norm,required
 assert all(t not in full for t in ['{{','证件照','[confirm]','18270896361','不重复计','不单独归因','本地配置模式','20,969','独立代码'])
 with ZipFile(docx) as z:assert (R/'assets/resume/headshot.png').read_bytes() in [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
 assert any(abs(im.image.width/im.image.height-.75)<.02 for im in reader.pages[0].images),'PDF headshot missing'
 target=R/'public'/profile['pdfPath'];target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(pdf.read_bytes())
 report={'pages':1,'a4':True,'text_searchable':len(full),'real_headshot':True,'data_sha256':hash_file(data_path),'pdf_sha256':hash_file(pdf),'template_sha256':hash_file(R/'assets/resume/resume_v2.docx'),'phone_public':False,'source':str(data_path),'visual_review':'pending'}
 (out/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');(R/'public/resume/manifest.json').write_text(json.dumps({k:v for k,v in report.items() if k not in ['source','visual_review']},ensure_ascii=False,indent=2)+'\n');print(out)
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output-dir',type=Path,default=R/'.resume-review/pdf');export(p.parse_args().output_dir.resolve())
